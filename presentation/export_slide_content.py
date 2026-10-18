"""Exports SLIDE_CONTENT.md from FINAL_PRESENTATION.pptx so that the two never diverge.

usage: <python with python-pptx> presentation/export_slide_content.py
"""

import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
DECK = HERE / "FINAL_PRESENTATION.pptx"
OUT = HERE / "SLIDE_CONTENT.md"

LAYOUT = {1: "Title band", 2: "Three cards", 3: "Numbered objectives + side panels", 4: "Full-width architecture diagram (native shapes)",
          5: "Three real n8n screenshots", 6: "Sequence of steps + latency cards", 7: "Table + technique chips", 8: "Split before/after + table",
          9: "Two before/after panels + two tables", 10: "Before/after + lesson + compliance table", 11: "Table + provenance flow + stat cards",
          12: "Pipeline diagram (native shapes)", 13: "Native column chart + table", 14: "Classification table + rationale",
          15: "Pipeline + results table + stat card", 16: "Four demo rows", 17: "Before/after stat cards + table", 18: "Two columns",
          19: "Four statement cards", 20: "Divider", 21: "Table + bullets", 22: "Table + native chart", 23: "Table + bullets", 24: "Table", 25: "Native chart"}
TRANSITION = {3: "Member 2 will now show you how the system is built.", 6: "That difference shaped the whole design (to Member 3).",
              10: "Member 4 explains how we find the articles.", 13: "Member 5 explains the agents that use these passages.",
              15: "Now Member 6 will show the system live."}


def slide_text(slide) -> list[str]:
    lines = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            t = " ".join(p.text.strip() for p in shp.text_frame.paragraphs if p.text.strip())
            if t:
                lines.append(t)
        elif shp.has_table:
            rows = [" | ".join(c.text for c in r.cells) for r in shp.table.rows]
            lines.append("TABLE: " + " // ".join(rows))
        elif shp.shape_type == 3 or getattr(shp, "has_chart", False):
            if getattr(shp, "has_chart", False):
                ch = shp.chart
                ser = "; ".join(f"{s.name}: {', '.join(f'{v:.3f}' for v in s.values)}" for s in ch.plots[0].series)
                cats = ", ".join(str(c) for c in ch.plots[0].categories)
                lines.append(f"CHART (column): categories [{cats}] — {ser}")
        elif shp.shape_type == 13:
            lines.append("PICTURE: real screenshot")
    return lines


def main() -> None:
    timing = json.loads(subprocess.run([sys.executable, str(HERE / "check_timing.py")], capture_output=True, text=True).stdout)
    secs = {r["slide"]: r["seconds"] for m in timing["members"] for r in m["slides"] if not r.get("note")}
    prs = Presentation(DECK)
    out = ["# Slide content (exported from FINAL_PRESENTATION.pptx by export_slide_content.py)", "",
           "Each block lists the slide's professor section, presenter, target duration, the exact on-slide text (tables, "
           "charts and pictures summarised), the main message (first line of the speaker notes), the evidence (the source "
           "line), the layout and the transition. Do not edit by hand; regenerate after changing the deck.", ""]
    for i, slide in enumerate(prs.slides, 1):
        lines = slide_text(slide)
        chip = next((l for l in lines if re.match(r"^\d · ", l)), "")
        presenter = "Member 1" if i == 1 else next((l for l in lines if l.startswith("Member") or l == "All members"), "")
        if i == 1:
            chip = "1 · Introduction and Use Case (title)"
        src = next((l for l in lines if l.startswith("Source: ")), "")
        body = [l for l in lines if l not in (chip, presenter, src) and not l.startswith("Tunisian Administrative Procedures Assistant ·") and not re.fullmatch(r"\d+", l)]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        message = notes.split("\n")[0][:300] if notes else ""
        out += [f"## Slide {i}", "",
                f"- **Section:** {chip or ('backup' if i > 19 else '—')}",
                f"- **Presenter:** {presenter or ('Member 1' if i == 1 else '—')}",
                f"- **Target duration:** {secs.get(i, 0)} s" + (" (+ demo time)" if i == 16 else "") if i <= 19 else "- **Target duration:** backup only",
                f"- **Layout:** {LAYOUT.get(i, '')}",
                f"- **Main message:** {message}",
                f"- **Evidence:** {src[8:] if src else 'see notes'}",
                f"- **Transition:** {TRANSITION.get(i, 'next slide, same presenter' if i < 19 else '—')}",
                "", "**On-slide text:**", ""]
        out += [f"- {l}" for l in body] + [""]
    OUT.write_text("\n".join(out), encoding="utf-8")
    print("written", OUT.name)


if __name__ == "__main__":
    main()
